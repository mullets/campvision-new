"""Planilha — o entregável da Fase 1.

A planilha é feita para REVISÃO, não só para arquivo: célula com confiança
baixa sai pintada, cada campo tem sua confiança numa coluna espelho, e a coluna
"OK?" é onde você marca o que já conferiu. A Fase 2 (aplicar) lê este mesmo
arquivo depois de você editar.
"""

from __future__ import annotations

import csv
import json
import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .esquema import CAMPOS, CAMPOS_POR_NOME, Leitura

_log = logging.getLogger("cv2.planilha")

LIMIAR_ATENCAO = 0.60
LIMIAR_BOM = 0.85

_VERMELHO = PatternFill("solid", fgColor="FFC7CE")
# Azul = veio da PASTA, não do carimbo. Cor própria de propósito: você tem que
# conseguir ver de relance o que foi lido e o que foi deduzido do caminho.
_AZUL = PatternFill("solid", fgColor="DDEBF7")
_AMARELO = PatternFill("solid", fgColor="FFEB9C")
_CINZA = PatternFill("solid", fgColor="EEEEEE")
_CABECALHO = PatternFill("solid", fgColor="1F3864")


def _colunas() -> list[str]:
    fixas = ["Arquivo", "OK?", "Grupo/Projeto", "Carimbo?", "Confiança média"]
    campos = [c.rotulo for c in CAMPOS]
    da_pasta = ["Fundo (pasta)", "Projeto (pasta)", "Ano (pasta)", "Divergência"]
    espelho = [f"conf. {c.rotulo}" for c in CAMPOS]
    lidos = [f"lido: {c.rotulo}" for c in CAMPOS if c.do_projeto]
    return (
        fixas + campos + da_pasta
        + ["Rotação", "Passes", "Nota da IA", "Erro"] + espelho + lidos
    )


def _linha(leitura: Leitura) -> list[object]:
    fixas: list[object] = [
        leitura.arquivo,
        "",
        leitura.grupo or leitura.valores.get("projeto", ""),
        "sim" if leitura.carimbo_encontrado else "NÃO",
        round(leitura.confianca_media, 2),
    ]
    campos = [leitura.valores.get(c.nome, "") for c in CAMPOS]
    rotulos_divergentes = ", ".join(
        CAMPOS_POR_NOME[n].rotulo for n in leitura.divergencias if n in CAMPOS_POR_NOME
    )
    da_pasta: list[object] = [
        leitura.pista_fundo,
        leitura.pista_projeto,
        leitura.pista_ano,
        f"carimbo ≠ pasta: {rotulos_divergentes}" if rotulos_divergentes else "",
    ]
    extras: list[object] = [
        leitura.rotacao,
        leitura.passes,
        leitura.nota_ia + (" [pode ser de outro projeto]" if leitura.suspeita_grupo else ""),
        leitura.erro,
    ]
    espelho = [round(leitura.confiancas.get(c.nome, 0.0), 2) for c in CAMPOS]
    lidos = [leitura.lidos_originais.get(c.nome, "") for c in CAMPOS if c.do_projeto]
    return fixas + campos + da_pasta + extras + espelho + lidos


def _ordenar(leituras: list[Leitura]) -> list[Leitura]:
    """Agrupadas por projeto, e dentro do projeto por número de folha."""

    def folha_num(leitura: Leitura) -> tuple[int, str]:
        bruto = (leitura.valores.get("folha") or "").strip()
        digitos = "".join(ch for ch in bruto.split("/")[0] if ch.isdigit())
        return (int(digitos), bruto) if digitos else (10**9, leitura.arquivo)

    return sorted(
        leituras,
        key=lambda l: ((l.grupo or l.valores.get("projeto") or "zzz").lower(), folha_num(l)),
    )


def escrever_xlsx(leituras: list[Leitura], destino: Path) -> Path:
    """Gera o XLSX de revisão. Devolve o caminho escrito."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Catalogação"
    colunas = _colunas()
    ws.append(colunas)

    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = _CABECALHO
        celula.alignment = Alignment(vertical="center", wrap_text=True)

    inicio_campos = 6  # 1-based: depois das 5 colunas fixas
    n_campos = len(CAMPOS)
    coluna_divergencia = inicio_campos + n_campos + 3
    inicio_espelho = inicio_campos + n_campos + 8

    for leitura in _ordenar(leituras):
        ws.append(_linha(leitura))
        linha = ws.max_row
        for i, campo in enumerate(CAMPOS):
            confianca = leitura.confiancas.get(campo.nome, 0.0)
            celula = ws.cell(row=linha, column=inicio_campos + i)
            if campo.nome in leitura.campos_da_pasta:
                celula.fill = _AZUL  # veio do caminho, não do carimbo
            elif not str(celula.value or "").strip():
                celula.fill = _CINZA
            elif confianca < LIMIAR_ATENCAO:
                celula.fill = _VERMELHO
            elif confianca < LIMIAR_BOM:
                celula.fill = _AMARELO
        if leitura.divergencias:
            ws.cell(row=linha, column=coluna_divergencia).fill = _VERMELHO
        if leitura.erro or not leitura.carimbo_encontrado:
            ws.cell(row=linha, column=4).fill = _VERMELHO

    ws.freeze_panes = "C2"
    for i, nome in enumerate(colunas, start=1):
        largura = 34 if i in (1, 3) else max(11, min(26, len(nome) + 4))
        ws.column_dimensions[get_column_letter(i)].width = largura
    # Colunas de auditoria ficam recolhidas: quem quiser confere, quem não quiser não vê.
    for col in range(inicio_espelho, len(colunas) + 1):
        ws.column_dimensions[get_column_letter(col)].hidden = True

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    _log.info("Planilha escrita: %s (%d linhas)", destino, len(leituras))
    return destino


def escrever_csv(leituras: list[Leitura], destino: Path) -> Path:
    destino.parent.mkdir(parents=True, exist_ok=True)
    with destino.open("w", encoding="utf-8-sig", newline="") as f:
        escritor = csv.writer(f)
        escritor.writerow(_colunas())
        for leitura in _ordenar(leituras):
            escritor.writerow(_linha(leitura))
    _log.info("CSV escrito: %s", destino)
    return destino


def escrever_json(leituras: list[Leitura], destino: Path) -> Path:
    """Leituras finais em JSON — a fonte para remontar a planilha do acervo."""
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text(
        json.dumps([l.para_dict() for l in leituras], ensure_ascii=False, indent=1),
        encoding="utf-8",
    )
    return destino


def ler_json(caminho: Path) -> list[Leitura]:
    """Lê o JSON de leituras. Arquivo corrompido devolve lista vazia com aviso."""
    try:
        dados = json.loads(caminho.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as erro:
        _log.warning("Não consegui ler %s: %s", caminho, erro)
        return []
    if not isinstance(dados, list):
        return []
    return [Leitura.de_dict(d) for d in dados if isinstance(d, dict)]


def escrever_relatorio(leituras: list[Leitura], destino: Path, custo_usd: float = 0.0) -> Path:
    """Resumo em texto: o que dá para olhar em 10 segundos e saber se o lote prestou."""
    total = len(leituras)
    com_carimbo = sum(1 for l in leituras if l.carimbo_encontrado)
    com_erro = sum(1 for l in leituras if l.erro)
    linhas = [
        "CAMP Vision 2 — relatório do lote",
        "=" * 40,
        f"Pranchas processadas: {total}",
        f"Carimbo encontrado:   {com_carimbo} ({com_carimbo / total * 100:.0f}%)" if total else "",
        f"Erros:                {com_erro}",
        f"Custo estimado:       US$ {custo_usd:.2f}",
        "",
        "Preenchimento por campo:",
    ]
    for campo in CAMPOS:
        n = sum(1 for l in leituras if l.valores.get(campo.nome))
        fracos = sum(
            1 for l in leituras
            if l.valores.get(campo.nome) and l.confiancas.get(campo.nome, 0) < LIMIAR_ATENCAO
        )
        pct = (n / total * 100) if total else 0
        linhas.append(f"  {campo.rotulo:<20} {n:>4}/{total} ({pct:>3.0f}%)  a revisar: {fracos}")

    grupos: dict[str, int] = {}
    for leitura in leituras:
        chave = leitura.grupo or leitura.valores.get("projeto") or "(sem projeto)"
        grupos[chave] = grupos.get(chave, 0) + 1
    linhas += ["", f"Projetos distintos: {len(grupos)}"]
    for nome, n in sorted(grupos.items(), key=lambda kv: -kv[1]):
        linhas.append(f"  {n:>4}  {nome}")

    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text("\n".join(l for l in linhas if l is not None), encoding="utf-8")
    return destino
