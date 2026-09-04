"""Fase 2 — aplicar a planilha revisada aos arquivos.

Roda SÓ depois de você revisar a planilha. Lê o XLSX/CSV editado e renomeia,
organiza em pastas e (opcionalmente) grava EXIF. Separar isto da leitura é o
ponto central do redesenho: leitura errada vira célula errada na planilha, que
você corrige em 5 segundos — nunca uma prancha perdida em pasta errada.

Modo simulação (padrão) só lista o que faria. Nada é apagado nunca: os
arquivos são COPIADOS para a pasta catalogada, o original fica onde está.
"""

from __future__ import annotations

import csv
import logging
import re
import shutil
import subprocess
from dataclasses import dataclass
from pathlib import Path

from .esquema import CAMPOS

_log = logging.getLogger("cv2.aplicar")

INVALIDOS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
ATRIBUICAO_PADRAO = "CAMP - Casa da Arquitetura Moderna Paulista"


@dataclass
class Acao:
    origem: Path
    destino: Path
    metadados: dict[str, str]


def limpar_nome(texto: str, limite: int = 80) -> str:
    """Deixa o texto usável como nome de arquivo/pasta."""
    limpo = INVALIDOS.sub("", (texto or "").strip())
    limpo = re.sub(r"\s+", " ", limpo).strip(" .")
    return limpo[:limite].strip() or "Sem titulo"


def ler_planilha(caminho: Path) -> list[dict[str, str]]:
    """Lê o XLSX ou CSV revisado e devolve as linhas como dicionários."""
    if caminho.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(caminho, data_only=True)
        ws = wb.active
        linhas = list(ws.iter_rows(values_only=True))
        if not linhas:
            return []
        cabecalho = [str(c or "") for c in linhas[0]]
        return [
            {cabecalho[i]: ("" if v is None else str(v)) for i, v in enumerate(linha) if i < len(cabecalho)}
            for linha in linhas[1:]
            if any(v is not None and str(v).strip() for v in linha)
        ]
    with caminho.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def montar_nome(linha: dict[str, str], sequencial: int) -> str:
    """`{Projeto} - {NNN} - {Título}`, sem depender de código automático.

    Simples de propósito: o nome do arquivo é derivado da planilha revisada, e
    a planilha é a fonte da verdade. Ajuste este método se sua convenção mudar.
    """
    projeto = limpar_nome(linha.get("Projeto", ""), 45)
    titulo = limpar_nome(linha.get("Título da prancha", "") or linha.get("Tipo", ""), 45)
    partes = [projeto, f"{sequencial:03d}"]
    if titulo and titulo != "Sem titulo":
        partes.append(titulo)
    return " - ".join(partes)


def montar_pasta(raiz: Path, linha: dict[str, str]) -> Path:
    ano = limpar_nome(linha.get("Ano", ""), 4) or "Ano desconhecido"
    projeto = limpar_nome(linha.get("Projeto", ""), 60) or "Sem projeto"
    return raiz / ano / projeto


def planejar(
    pasta_origem: Path, planilha: Path, raiz_destino: Path | None = None
) -> list[Acao]:
    """Monta a lista de ações sem tocar em disco."""
    linhas = ler_planilha(planilha)
    raiz = raiz_destino or pasta_origem.parent / f"{pasta_origem.name}_catalogado"
    contador: dict[str, int] = {}
    acoes: list[Acao] = []

    for linha in linhas:
        nome_arquivo = (linha.get("Arquivo") or "").strip()
        if not nome_arquivo:
            continue
        origem = pasta_origem / nome_arquivo
        if not origem.exists():
            _log.warning("Arquivo da planilha não existe na pasta: %s", nome_arquivo)
            continue
        pasta = montar_pasta(raiz, linha)
        chave = str(pasta)
        contador[chave] = contador.get(chave, 0) + 1
        destino = pasta / f"{montar_nome(linha, contador[chave])}{origem.suffix}"
        metadados = {c.rotulo: linha.get(c.rotulo, "") for c in CAMPOS}
        acoes.append(Acao(origem=origem, destino=destino, metadados=metadados))

    return acoes


def executar(acoes: list[Acao], simular: bool = True, gravar_exif: bool = False) -> list[str]:
    """Copia os arquivos para o destino. `simular=True` só descreve."""
    mensagens: list[str] = []
    for acao in acoes:
        if simular:
            mensagens.append(f"[simulação] {acao.origem.name} -> {acao.destino}")
            continue
        acao.destino.parent.mkdir(parents=True, exist_ok=True)
        final = acao.destino
        n = 2
        while final.exists():  # nunca sobrescreve
            final = acao.destino.with_stem(f"{acao.destino.stem} ({n})")
            n += 1
        shutil.copy2(acao.origem, final)
        mensagens.append(f"{acao.origem.name} -> {final}")
        if gravar_exif:
            erro = _gravar_exif(final, acao.metadados)
            if erro:
                mensagens.append(f"  aviso EXIF: {erro}")
    return mensagens


def _gravar_exif(caminho: Path, metadados: dict[str, str], atribuicao: str = ATRIBUICAO_PADRAO) -> str:
    """Grava metadados via exiftool, se disponível. Copyright = autor + instituição."""
    if not shutil.which("exiftool"):
        return "exiftool não encontrado no PATH — pulei os metadados."
    autor = metadados.get("Arquiteto") or metadados.get("Escritório") or ""
    copyright_ = f"{autor} / {atribuicao}" if autor else atribuicao
    descricao = " | ".join(
        f"{rotulo}: {valor}" for rotulo, valor in metadados.items() if valor
    )[:1800]
    argumentos = [
        "exiftool", "-overwrite_original", "-charset", "utf8",
        f"-Artist={autor}",
        f"-Copyright={copyright_}",
        f"-XMP-dc:Rights={copyright_}",
        f"-XMP-dc:Title={metadados.get('Título da prancha', '')}",
        f"-ImageDescription={descricao}",
        str(caminho),
    ]
    try:
        proc = subprocess.run(argumentos, capture_output=True, text=True, timeout=90)
    except (OSError, subprocess.TimeoutExpired) as erro:
        return str(erro)
    return "" if proc.returncode == 0 else proc.stderr.strip()[:200]
