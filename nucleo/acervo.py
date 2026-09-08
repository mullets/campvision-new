"""Planilha única do acervo.

Junta o que cada projeto produziu numa planilha só, na raiz. Não gasta API:
é remontada dos `catalogacao/leituras.json` de cada projeto, então você pode
rodar quantas vezes quiser, a qualquer momento, e ela sempre reflete o estado
atual do acervo.

Três abas:
  Acervo   — uma linha por prancha, o acervo inteiro
  Projetos — uma linha por projeto: quantas pranchas, cobertura, o que revisar
  Pendentes— o que ainda não passou pela leitura
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import planilha as mod_planilha
from .config import Config
from .esquema import CAMPOS, Leitura
from .planilha import LIMIAR_ATENCAO, _VERMELHO, _AMARELO, _CINZA, _CABECALHO
from .vigia import NOME_STATUS, descobrir, ler_status

_log = logging.getLogger("cv2.acervo")

NOME_ARQUIVO = "acervo.xlsx"


@dataclass
class ProjetoNoAcervo:
    nome: str
    caminho_relativo: str
    leituras: list[Leitura]
    status: str
    fase: str

    @property
    def pranchas(self) -> int:
        return len(self.leituras)

    @property
    def com_carimbo(self) -> int:
        return sum(1 for l in self.leituras if l.carimbo_encontrado)

    @property
    def a_revisar(self) -> int:
        return sum(
            1
            for l in self.leituras
            for nome, valor in l.valores.items()
            if valor and l.confiancas.get(nome, 0.0) < LIMIAR_ATENCAO
        )

    @property
    def fundo(self) -> str:
        for leitura in self.leituras:
            if leitura.pista_fundo:
                return leitura.pista_fundo
        return ""

    @property
    def divergencias(self) -> int:
        return sum(1 for l in self.leituras if l.divergencias)

    @property
    def anos(self) -> str:
        valores = sorted({l.valores.get("ano", "") for l in self.leituras if l.valores.get("ano")})
        if not valores:
            return ""
        return valores[0] if len(valores) == 1 else f"{valores[0]}–{valores[-1]}"

    def campo_predominante(self, campo: str) -> str:
        contagem: dict[str, int] = {}
        for leitura in self.leituras:
            valor = (leitura.valores.get(campo) or "").strip()
            if valor:
                contagem[valor] = contagem.get(valor, 0) + 1
        return max(contagem, key=contagem.get) if contagem else ""


def estimar(raiz: Path, config: Config, historico: list | None = None) -> dict:
    """Conta o que seria processado e estima custo e tempo, sem chamar a API.

    Números aproximados de propósito: servem para decidir se vale soltar o
    mutirão, não para fechar orçamento. O custo real aparece ao vivo no painel.
    """
    from .vigia import varrer

    pendentes = varrer(raiz, config)
    pranchas = 0
    por_projeto: list[tuple[str, int]] = []
    for projeto in pendentes:
        n = len(projeto.arquivos(config))
        pranchas += n
        try:
            rel = str(projeto.pasta.relative_to(raiz))
        except ValueError:
            rel = projeto.pasta.name
        por_projeto.append((rel, n))

    # Se já há histórico de lotes reais, o custo médio por prancha medido vale
    # muito mais que qualquer conta minha — o primeiro palpite estava pela
    # metade do observado. Sem histórico, cai no valor calibrado abaixo.
    medido = _custo_medio_por_prancha(historico or [])
    if medido:
        piso, teto = pranchas * medido * 0.8, pranchas * medido * 1.5
        origem = f"medido no seu acervo (US$ {medido:.3f}/prancha)"
    else:
        # Calibrado contra lotes reais: ~US$ 0,02 por chamada. O piso supõe 1
        # chamada por prancha (cache de região funcionando); o teto, 2.
        entrada, saida = 3500, 1000
        piso = config.custo_estimado_usd(pranchas * entrada, pranchas * saida)
        teto = config.custo_estimado_usd(pranchas * entrada * 2, pranchas * saida * 2)
        origem = "estimativa (sem histórico ainda)"

    # ~4s por chamada, dividido pelos trabalhadores.
    segundos = pranchas * 4 / max(1, config.trabalhadores)

    return {
        "projetos": len(pendentes),
        "pranchas": pranchas,
        "origem": origem,
        "custo_min": piso,
        "custo_max": teto,
        "horas": segundos / 3600,
        "por_projeto": sorted(por_projeto, key=lambda kv: -kv[1]),
    }


def _custo_medio_por_prancha(eventos: list) -> float:
    """Custo por prancha observado nos lotes já rodados. 0 se não há amostra."""
    pranchas = sum(e.pranchas for e in eventos if not e.falha)
    custo = sum(e.custo_usd for e in eventos if not e.falha)
    if pranchas < 20 or custo <= 0:  # amostra pequena demais para confiar
        return 0.0
    return custo / pranchas


def _ler_fase(caminho: Path) -> str:
    import json

    if not caminho.exists():
        return ""
    try:
        dados = json.loads(caminho.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return ""
    return str(dados.get("fase", "")) if isinstance(dados, dict) else ""


def levantar(raiz: Path, config: Config) -> tuple[list[ProjetoNoAcervo], list[str]]:
    """Percorre o acervo e devolve (processados, pendentes)."""
    processados: list[ProjetoNoAcervo] = []
    pendentes: list[str] = []

    for projeto in descobrir(raiz, config):
        try:
            relativo = str(projeto.pasta.relative_to(raiz))
        except ValueError:
            relativo = projeto.pasta.name
        json_leituras = projeto.pasta / "catalogacao" / "leituras.json"
        if not json_leituras.exists():
            pendentes.append(relativo)
            continue
        leituras = mod_planilha.ler_json(json_leituras)
        if not leituras:
            pendentes.append(relativo)
            continue
        processados.append(
            ProjetoNoAcervo(
                nome=projeto.nome,
                caminho_relativo=relativo,
                leituras=leituras,
                status=ler_status(projeto.pasta / NOME_STATUS),
                fase=_ler_fase(projeto.pasta / NOME_STATUS),
            )
        )
    return processados, pendentes


def _aba_acervo(wb: Workbook, projetos: list[ProjetoNoAcervo]) -> None:
    ws = wb.active
    ws.title = "Acervo"
    colunas = ["Pasta", "Arquivo", "OK?", "Carimbo?", "Confiança"] + [c.rotulo for c in CAMPOS]
    ws.append(colunas)
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = _CABECALHO
        celula.alignment = Alignment(vertical="center", wrap_text=True)

    inicio_campos = 6
    for projeto in sorted(projetos, key=lambda p: p.caminho_relativo.lower()):
        for leitura in projeto.leituras:
            ws.append(
                [
                    projeto.caminho_relativo,
                    leitura.arquivo,
                    "",
                    "sim" if leitura.carimbo_encontrado else "NÃO",
                    round(leitura.confianca_media, 2),
                ]
                + [leitura.valores.get(c.nome, "") for c in CAMPOS]
            )
            linha = ws.max_row
            for i, campo in enumerate(CAMPOS):
                celula = ws.cell(row=linha, column=inicio_campos + i)
                confianca = leitura.confiancas.get(campo.nome, 0.0)
                if not str(celula.value or "").strip():
                    celula.fill = _CINZA
                elif confianca < LIMIAR_ATENCAO:
                    celula.fill = _VERMELHO
                elif confianca < 0.85:
                    celula.fill = _AMARELO

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    for i, nome in enumerate(colunas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (
            38 if i == 1 else 30 if i == 2 else max(11, min(26, len(nome) + 4))
        )


def _aba_projetos(wb: Workbook, projetos: list[ProjetoNoAcervo]) -> None:
    ws = wb.create_sheet("Projetos")
    colunas = [
        "Pasta", "Fundo", "Projeto (lido)", "Arquiteto", "Cidade", "Ano(s)",
        "Pranchas", "Com carimbo", "%", "A revisar", "Divergências", "Status", "Fase",
    ]
    ws.append(colunas)
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = _CABECALHO

    for projeto in sorted(projetos, key=lambda p: -p.pranchas):
        pct = round(projeto.com_carimbo / projeto.pranchas * 100) if projeto.pranchas else 0
        ws.append([
            projeto.caminho_relativo,
            projeto.fundo,
            projeto.campo_predominante("projeto"),
            projeto.campo_predominante("arquiteto"),
            projeto.campo_predominante("cidade"),
            projeto.anos,
            projeto.pranchas,
            projeto.com_carimbo,
            pct,
            projeto.a_revisar,
            projeto.divergencias,
            projeto.status,
            projeto.fase,
        ])
        if pct < 80:
            ws.cell(row=ws.max_row, column=9).fill = _VERMELHO
        elif pct < 95:
            ws.cell(row=ws.max_row, column=9).fill = _AMARELO
        if projeto.divergencias:
            ws.cell(row=ws.max_row, column=11).fill = _VERMELHO

    total = sum(p.pranchas for p in projetos)
    carimbos = sum(p.com_carimbo for p in projetos)
    ws.append([])
    ws.append([
        f"TOTAL: {len(projetos)} projeto(s)", "", "", "", "", "",
        total, carimbos,
        round(carimbos / total * 100) if total else 0,
        sum(p.a_revisar for p in projetos),
        sum(p.divergencias for p in projetos), "", "",
    ])
    for celula in ws[ws.max_row]:
        celula.font = Font(bold=True)

    ws.freeze_panes = "B2"
    for i, nome in enumerate(colunas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 38 if i in (1, 3) else max(11, len(nome) + 4)


def _aba_pendentes(wb: Workbook, pendentes: list[str]) -> None:
    ws = wb.create_sheet("Pendentes")
    ws.append(["Pasta ainda não catalogada"])
    ws["A1"].font = Font(bold=True, color="FFFFFF")
    ws["A1"].fill = _CABECALHO
    for caminho in sorted(pendentes):
        ws.append([caminho])
    if not pendentes:
        ws.append(["(nenhuma — o acervo inteiro passou pela leitura)"])
    ws.column_dimensions["A"].width = 70


def escrever(raiz: Path, config: Config) -> tuple[Path, int, int]:
    """Monta a planilha única. Devolve (caminho, projetos, pranchas)."""
    projetos, pendentes = levantar(raiz, config)
    wb = Workbook()
    _aba_acervo(wb, projetos)
    _aba_projetos(wb, projetos)
    _aba_pendentes(wb, pendentes)

    destino = raiz / config.pasta_acervo / NOME_ARQUIVO
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)

    # CSV chapado junto, para quem prefere grep a Excel.
    todas = [l for p in projetos for l in p.leituras]
    mod_planilha.escrever_csv(todas, destino.with_name("acervo.csv"))

    pranchas = sum(p.pranchas for p in projetos)
    _log.info(
        "Planilha do acervo: %d projeto(s), %d prancha(s), %d pendente(s) → %s",
        len(projetos), pranchas, len(pendentes), destino,
    )
    return destino, len(projetos), pranchas
